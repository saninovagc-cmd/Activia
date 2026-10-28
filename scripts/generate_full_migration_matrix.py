import os
import openpyxl
import json

base_dir = r"c:\Users\LENOVO X13\Desktop\ITA\Activia\Base de données"
matrix = []

for root, dirs, files in os.walk(base_dir):
    service = os.path.basename(root)
    if service == "Base de données":
        continue
    for f in sorted(files):
        if not f.endswith(".xlsx"):
            continue
        p = os.path.join(root, f)
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            for sname in wb.sheetnames:
                sh = wb[sname]
                # count rows that have actual data
                count = 0
                first_row = None
                for row in sh.iter_rows(values_only=True):
                    non_empty = [c for c in row if c is not None and str(c).strip() != ""]
                    if non_empty:
                        count += 1
                        if first_row is None and len(non_empty) >= 3:
                            first_row = [str(c).replace('\n', ' ').strip() for c in non_empty[:8]]
                
                matrix.append({
                    "service": service,
                    "file": f,
                    "sheet": sname,
                    "active_rows": count,
                    "headers_preview": first_row or []
                })
            wb.close()
        except Exception as e:
            matrix.append({
                "service": service,
                "file": f,
                "error": str(e)
            })

with open(r"c:\Users\LENOVO X13\Desktop\ITA\Activia\scripts\migration_matrix.json", "w", encoding="utf-8") as out:
    json.dump(matrix, out, ensure_ascii=False, indent=2)

print(f"Matrix computed for {len(matrix)} sheets across all Excel files.")
