import openpyxl, os

print("=== INSPECTING PHASE 4 FILES ===")

# 1. SSMUR Signalements
f_sig = r'Base de données/SSMUR/Base_Suivi_Signalements_Produits_Sante_.xlsx'
if os.path.exists(f_sig):
    wb = openpyxl.load_workbook(f_sig, read_only=True, data_only=True)
    print(f"\nFile: {os.path.basename(f_sig)} -> Sheets: {wb.sheetnames}")
    for sname in wb.sheetnames:
        sh = wb[sname]
        r0 = next(sh.iter_rows(values_only=True, max_col=20), None)
        cnt = 0
        empty = 0
        for r in sh.iter_rows(values_only=True, max_col=20):
            if any(r[:5]):
                cnt += 1
            else:
                empty += 1
                if empty > 15:
                    break
        print(f"  [{sname}]: {cnt} rows. Header: {[str(c)[:30] for c in r0 if c][:8] if r0 else None}")
    wb.close()

# 2. SSMUR Echantillonnage
f_ech = r'Base de données/SSMUR/BASE DE DONEES ABMED_ECHANTILLONNAGE_BN_2026 PHASE _COMPILATION 23092026 (5).xlsx'
if os.path.exists(f_ech):
    wb = openpyxl.load_workbook(f_ech, read_only=True, data_only=True)
    print(f"\nFile: {os.path.basename(f_ech)} -> Sheets: {wb.sheetnames}")
    for sname in wb.sheetnames:
        sh = wb[sname]
        r0 = next(sh.iter_rows(values_only=True, max_col=20), None)
        cnt = 0
        empty = 0
        for r in sh.iter_rows(values_only=True, max_col=20):
            if any(r[:5]):
                cnt += 1
            else:
                empty += 1
                if empty > 15:
                    break
        print(f"  [{sname}]: {cnt} rows. Header: {[str(c)[:30] for c in r0 if c][:8] if r0 else None}")
    wb.close()

# 3. SVPS Formations
f_form = r'Base de données/SVPS/19_Registre de formation des points focaux.xlsx'
if os.path.exists(f_form):
    wb = openpyxl.load_workbook(f_form, read_only=True, data_only=True)
    print(f"\nFile: {os.path.basename(f_form)} -> Sheets: {wb.sheetnames}")
    for sname in wb.sheetnames:
        sh = wb[sname]
        r0 = next(sh.iter_rows(values_only=True, max_col=20), None)
        cnt = 0
        empty = 0
        for r in sh.iter_rows(values_only=True, max_col=20):
            if any(r[:5]):
                cnt += 1
            else:
                empty += 1
                if empty > 15:
                    break
        print(f"  [{sname}]: {cnt} rows. Header: {[str(c)[:30] for c in r0 if c][:8] if r0 else None}")
    wb.close()

# 4. SVPS MAPI files
for mapi_f in ['Campagne HPV_2025 Suivi imputabilit cas MAPI.xlsx', 'Campagne nVPO2_2025 Suivi imputabilit cas MAPI.xlsx', 'Campagne CPS_2026 Suivi imputabilit.xlsx']:
    fpath = os.path.join(r'Base de données/SVPS', mapi_f)
    if not os.path.exists(fpath):
        # find matching file
        for real_f in os.listdir(r'Base de données/SVPS'):
            if real_f.startswith(mapi_f[:12]):
                fpath = os.path.join(r'Base de données/SVPS', real_f)
                break
    if os.path.exists(fpath):
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        print(f"\nFile: {os.path.basename(fpath)} -> Sheets: {wb.sheetnames}")
        sh = wb[wb.sheetnames[0]]
        r0 = next(sh.iter_rows(values_only=True, max_col=20), None)
        cnt = 0
        empty = 0
        for r in sh.iter_rows(values_only=True, max_col=20):
            if any(r[:5]):
                cnt += 1
            else:
                empty += 1
                if empty > 15:
                    break
        print(f"  [{wb.sheetnames[0]}]: {cnt} rows. Header: {[str(c)[:30] for c in r0 if c][:8] if r0 else None}")
        wb.close()
