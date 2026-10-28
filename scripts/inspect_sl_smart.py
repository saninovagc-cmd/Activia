import openpyxl, os

sl_dir = r'Base de données/SL'
print('Files in SL directory:')
for f in os.listdir(sl_dir):
    fpath = os.path.join(sl_dir, f)
    sz = os.path.getsize(fpath) / 1024
    print(f'  {f} ({sz:.1f} KB)')

wb_path = os.path.join(sl_dir, 'Activités du SL.xlsx')
if os.path.exists(wb_path):
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    print('\nSheets in Activités du SL.xlsx:', wb.sheetnames)
    for sname in wb.sheetnames:
        sh = wb[sname]
        headers = None
        count = 0
        empty_streak = 0
        sample_rows = []
        for r in sh.iter_rows(values_only=True):
            non_empty = [str(c).replace('\n', ' ').strip() if c is not None else '' for c in r]
            has_val = any(c for c in non_empty[:8])
            if headers is None:
                if len([c for c in non_empty if c]) >= 3:
                    headers = non_empty
                continue
            if not has_val:
                empty_streak += 1
                if empty_streak > 20:
                    break
                continue
            empty_streak = 0
            count += 1
            if len(sample_rows) < 3:
                sample_rows.append(non_empty)
        
        clean_hdrs = [h for h in headers if h][:8] if headers else []
        print(f"\nSheet [{sname}]: {count} data rows")
        print(f"  Headers: {clean_hdrs}")
        for idx, s in enumerate(sample_rows):
            print(f"  Sample {idx+1}: {[c[:30] for c in s if c][:6]}")
    wb.close()
