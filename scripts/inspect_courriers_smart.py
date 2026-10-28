import openpyxl

wb = openpyxl.load_workbook(r'Base de données/DLVS/TABLEAU DE BORD COURRIERS DLVS_15.09.2026.xlsx', read_only=True, data_only=True)
print("Sheets:", wb.sheetnames)

for sname in ['SSMUR', 'SL', 'SVPS']:
    sh = wb[sname]
    print(f"\n================ SHEET: {sname} ================")
    headers = None
    count = 0
    empty_streak = 0
    sample_rows = []

    for row in sh.iter_rows(values_only=True):
        non_empty = [str(c).replace('\n', ' ').strip() if c is not None else '' for c in row]
        has_val = any(c for c in non_empty[:10])
        
        if headers is None:
            if len([c for c in non_empty if c]) >= 3:
                headers = non_empty
                print(f"Header: {headers[:15]} ({len(headers)} cols)")
            continue

        if not has_val:
            empty_streak += 1
            if empty_streak > 30:
                break
            continue
        
        empty_streak = 0
        count += 1
        if len(sample_rows) < 4:
            sample_rows.append(non_empty)

    print(f"Actual data rows in {sname}: {count}")
    for idx, r in enumerate(sample_rows):
        print(f"  Sample {idx+1}: {r[:12]}")

wb.close()
